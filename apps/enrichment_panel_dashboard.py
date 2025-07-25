# enrichment_panel_dashboard.py
import panel as pn
import pandas as pd
import numpy as np
from datetime import datetime, timedelta, date
import hashlib
import json
import os
from functools import wraps
import param
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Enable Panel extensions
pn.extension('tabulator', 'bokeh', notifications=True)

# Apple-style CSS
APPLE_CSS = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=SF+Pro+Display:wght@300;400;500;600;700&display=swap');
    
    * {
        font-family: -apple-system, BlinkMacSystemFont, 'SF Pro Display', 'Segoe UI', 'Roboto', sans-serif;
    }
    
    body {
        background-color: #f5f5f7;
        margin: 0;
        padding: 0;
    }
    
    .main-container {
        max-width: 1200px;
        margin: 0 auto;
        padding: 20px;
    }
    
    .card {
        background: rgba(255, 255, 255, 0.9);
        backdrop-filter: blur(20px);
        border-radius: 16px;
        padding: 24px;
        margin-bottom: 20px;
        box-shadow: 0 4px 20px rgba(0,0,0,0.08);
        border: 1px solid rgba(0,0,0,0.05);
        transition: all 0.3s ease;
    }
    
    .card:hover {
        box-shadow: 0 8px 30px rgba(0,0,0,0.12);
        transform: translateY(-2px);
    }
    
    .header {
        background: linear-gradient(135deg, #007AFF 0%, #5856D6 100%);
        color: white;
        padding: 40px;
        border-radius: 20px;
        margin-bottom: 30px;
        text-align: center;
    }
    
    .header h1 {
        font-size: 48px;
        font-weight: 700;
        margin: 0;
        letter-spacing: -2px;
    }
    
    .header p {
        font-size: 20px;
        font-weight: 300;
        margin-top: 10px;
        opacity: 0.9;
    }
    
    .stat-card {
        background: linear-gradient(135deg, #FF3B30 0%, #FF6482 100%);
        color: white;
        padding: 30px;
        border-radius: 16px;
        text-align: center;
        transition: all 0.3s ease;
    }
    
    .stat-card.blue {
        background: linear-gradient(135deg, #007AFF 0%, #5AC8FA 100%);
    }
    
    .stat-card.green {
        background: linear-gradient(135deg, #34C759 0%, #A3E048 100%);
    }
    
    .stat-card.purple {
        background: linear-gradient(135deg, #5856D6 0%, #AF52DE 100%);
    }
    
    .stat-value {
        font-size: 48px;
        font-weight: 700;
        margin: 10px 0;
    }
    
    .stat-label {
        font-size: 16px;
        font-weight: 500;
        opacity: 0.9;
    }
    
    .button {
        background: #007AFF;
        color: white;
        border: none;
        padding: 12px 24px;
        border-radius: 12px;
        font-size: 16px;
        font-weight: 600;
        cursor: pointer;
        transition: all 0.3s ease;
        display: inline-block;
        text-decoration: none;
    }
    
    .button:hover {
        background: #0051D5;
        transform: translateY(-1px);
        box-shadow: 0 4px 15px rgba(0,122,255,0.3);
    }
    
    .button.secondary {
        background: #F2F2F7;
        color: #000;
    }
    
    .button.secondary:hover {
        background: #E5E5EA;
    }
    
    .button.danger {
        background: #FF3B30;
    }
    
    .button.danger:hover {
        background: #D70015;
    }
    
    .input-field {
        width: 100%;
        padding: 14px 18px;
        border: 1px solid #E5E5EA;
        border-radius: 10px;
        font-size: 16px;
        transition: all 0.3s ease;
        background: #FFFFFF;
    }
    
    .input-field:focus {
        outline: none;
        border-color: #007AFF;
        box-shadow: 0 0 0 3px rgba(0,122,255,0.1);
    }
    
    .label {
        font-size: 14px;
        font-weight: 600;
        color: #1D1D1F;
        margin-bottom: 8px;
        display: block;
    }
    
    .table-container {
        overflow-x: auto;
        border-radius: 12px;
        border: 1px solid #E5E5EA;
    }
    
    table {
        width: 100%;
        border-collapse: collapse;
        background: white;
    }
    
    th {
        background: #F2F2F7;
        padding: 16px;
        text-align: left;
        font-weight: 600;
        font-size: 14px;
        color: #1D1D1F;
        border-bottom: 1px solid #E5E5EA;
    }
    
    td {
        padding: 16px;
        border-bottom: 1px solid #F2F2F7;
        font-size: 15px;
        color: #3A3A3C;
    }
    
    tr:hover {
        background: #FAFAFA;
    }
    
    .nav-tabs {
        display: flex;
        gap: 10px;
        margin-bottom: 30px;
        background: #F2F2F7;
        padding: 4px;
        border-radius: 12px;
    }
    
    .nav-tab {
        padding: 10px 20px;
        border-radius: 8px;
        font-weight: 600;
        cursor: pointer;
        transition: all 0.3s ease;
        color: #3A3A3C;
    }
    
    .nav-tab.active {
        background: white;
        color: #007AFF;
        box-shadow: 0 2px 8px rgba(0,0,0,0.1);
    }
    
    .progress-bar {
        width: 100%;
        height: 8px;
        background: #E5E5EA;
        border-radius: 4px;
        overflow: hidden;
    }
    
    .progress-fill {
        height: 100%;
        background: linear-gradient(90deg, #007AFF 0%, #5AC8FA 100%);
        transition: width 0.3s ease;
    }
    
    .alert {
        padding: 16px 20px;
        border-radius: 12px;
        margin-bottom: 20px;
        font-weight: 500;
    }
    
    .alert.success {
        background: #D1F4D1;
        color: #00611C;
    }
    
    .alert.error {
        background: #FFD1D1;
        color: #A80000;
    }
    
    .alert.info {
        background: #D1E7FF;
        color: #004AAD;
    }
    
    .loading {
        display: inline-block;
        width: 20px;
        height: 20px;
        border: 3px solid rgba(0,122,255,0.3);
        border-radius: 50%;
        border-top-color: #007AFF;
        animation: spin 1s ease-in-out infinite;
    }
    
    @keyframes spin {
        to { transform: rotate(360deg); }
    }
    
    .empty-state {
        text-align: center;
        padding: 60px 20px;
        color: #8E8E93;
    }
    
    .empty-state h3 {
        font-size: 24px;
        font-weight: 600;
        margin-bottom: 10px;
        color: #3A3A3C;
    }
    
    .modal-overlay {
        position: fixed;
        top: 0;
        left: 0;
        right: 0;
        bottom: 0;
        background: rgba(0,0,0,0.5);
        display: flex;
        align-items: center;
        justify-content: center;
        z-index: 1000;
    }
    
    .modal {
        background: white;
        border-radius: 20px;
        padding: 40px;
        max-width: 500px;
        width: 90%;
        box-shadow: 0 20px 60px rgba(0,0,0,0.3);
    }
</style>
"""

# Excel file configuration
EXCEL_FILE = "enrichment_data.xlsx"
SHEETS = {
    'users': 'Users',
    'schools': 'Schools',
    'cohorts': 'Cohorts',
    'courses': 'Courses',
    'assignments': 'Assignments',
    'sat_tests': 'SAT_Tests',
    'test_attempts': 'Test_Attempts',
    'surveys': 'Surveys',
    'survey_responses': 'Survey_Responses',
    'tutor_logs': 'Tutor_Logs'
}

# Initialize Excel file if it doesn't exist
def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        
        # Users sheet
        ws = wb.active
        ws.title = SHEETS['users']
        ws.append(['id', 'username', 'email', 'password_hash', 'role', 'school_id', 'cohort_id', 'created_at', 'is_active'])
        
        # Schools sheet
        ws = wb.create_sheet(SHEETS['schools'])
        ws.append(['id', 'name', 'address', 'contact_email'])
        
        # Cohorts sheet
        ws = wb.create_sheet(SHEETS['cohorts'])
        ws.append(['id', 'name', 'school_id', 'start_date', 'end_date', 'tutor_id'])
        
        # Courses sheet
        ws = wb.create_sheet(SHEETS['courses'])
        ws.append(['id', 'title', 'description', 'course_type', 'cohort_id'])
        
        # Assignments sheet
        ws = wb.create_sheet(SHEETS['assignments'])
        ws.append(['id', 'title', 'description', 'course_id', 'due_date', 'created_at'])
        
        # SAT Tests sheet
        ws = wb.create_sheet(SHEETS['sat_tests'])
        ws.append(['id', 'test_name', 'questions_data', 'time_limit', 'created_at'])
        
        # Test Attempts sheet
        ws = wb.create_sheet(SHEETS['test_attempts'])
        ws.append(['id', 'student_id', 'test_id', 'answers', 'score', 'started_at', 'completed_at'])
        
        # Surveys sheet
        ws = wb.create_sheet(SHEETS['surveys'])
        ws.append(['id', 'title', 'survey_type', 'questions', 'cohort_id'])
        
        # Survey Responses sheet
        ws = wb.create_sheet(SHEETS['survey_responses'])
        ws.append(['id', 'student_id', 'survey_id', 'responses', 'completed_at'])
        
        # Tutor Logs sheet
        ws = wb.create_sheet(SHEETS['tutor_logs'])
        ws.append(['id', 'tutor_id', 'date', 'hours_worked', 'activity_description', 'location', 'cohort_id', 'created_at'])
        
        # Style headers
        for sheet in wb.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="E5E5EA", end_color="E5E5EA", fill_type="solid")
        
        wb.save(EXCEL_FILE)
        
        # Create sample data
        create_sample_data()

# Excel helper functions
def read_sheet(sheet_name):
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=sheet_name)
        return df
    except:
        return pd.DataFrame()

def write_sheet(df, sheet_name):
    with pd.ExcelWriter(EXCEL_FILE, mode='a', if_sheet_exists='replace') as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

def append_to_sheet(data_dict, sheet_name):
    df = read_sheet(sheet_name)
    new_id = df['id'].max() + 1 if len(df) > 0 else 1
    data_dict['id'] = new_id
    df = pd.concat([df, pd.DataFrame([data_dict])], ignore_index=True)
    write_sheet(df, sheet_name)
    return new_id

def update_in_sheet(sheet_name, id_value, updates):
    df = read_sheet(sheet_name)
    mask = df['id'] == id_value
    for col, value in updates.items():
        df.loc[mask, col] = value
    write_sheet(df, sheet_name)

# Authentication functions
def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def verify_password(password, password_hash):
    return hash_password(password) == password_hash

# Session management
class Session:
    def __init__(self):
        self.user = None
        self.user_id = None
        self.role = None
        
    def login(self, username, password):
        users_df = read_sheet(SHEETS['users'])
        user = users_df[users_df['username'] == username]
        
        if len(user) > 0 and verify_password(password, user.iloc[0]['password_hash']) and user.iloc[0]['is_active']:
            self.user = user.iloc[0].to_dict()
            self.user_id = self.user['id']
            self.role = self.user['role']
            return True
        return False
    
    def logout(self):
        self.user = None
        self.user_id = None
        self.role = None
        
    def is_authenticated(self):
        return self.user is not None
    
    def has_role(self, roles):
        if isinstance(roles, str):
            roles = [roles]
        return self.role in roles

# Global session
session = Session()

# Create sample data
def create_sample_data():
    # Clear existing data
    wb = openpyxl.load_workbook(EXCEL_FILE)
    for sheet_name in SHEETS.values():
        ws = wb[sheet_name]
        ws.delete_rows(2, ws.max_row)
    wb.save(EXCEL_FILE)
    
    # Create sample school
    school_id = append_to_sheet({
        'name': 'Lincoln High School',
        'address': '123 Education St, Springfield, IL 62701',
        'contact_email': 'admin@lincolnhigh.edu'
    }, SHEETS['schools'])
    
    # Create sample cohort
    cohort_id = append_to_sheet({
        'name': 'Spring 2025 SAT Prep Cohort',
        'school_id': school_id,
        'start_date': date(2025, 1, 15).isoformat(),
        'end_date': date(2025, 5, 15).isoformat(),
        'tutor_id': 0  # Will update later
    }, SHEETS['cohorts'])
    
    # Create sample users
    admin_id = append_to_sheet({
        'username': 'admin',
        'email': 'admin@enrichment.com',
        'password_hash': hash_password('admin123'),
        'role': 'admin',
        'school_id': 0,
        'cohort_id': 0,
        'created_at': datetime.now().isoformat(),
        'is_active': True
    }, SHEETS['users'])
    
    tutor_id = append_to_sheet({
        'username': 'tutor1',
        'email': 'tutor@enrichment.com',
        'password_hash': hash_password('tutor123'),
        'role': 'tutor',
        'school_id': school_id,
        'cohort_id': 0,
        'created_at': datetime.now().isoformat(),
        'is_active': True
    }, SHEETS['users'])
    
    # Update cohort with tutor
    update_in_sheet(SHEETS['cohorts'], cohort_id, {'tutor_id': tutor_id})
    
    # Create sample students
    student_ids = []
    for name in ['alice', 'bob', 'carol']:
        student_id = append_to_sheet({
            'username': name,
            'email': f'{name}@student.edu',
            'password_hash': hash_password('student123'),
            'role': 'student',
            'school_id': school_id,
            'cohort_id': cohort_id,
            'created_at': datetime.now().isoformat(),
            'is_active': True
        }, SHEETS['users'])
        student_ids.append(student_id)
    
    # Create sample courses
    course_ids = []
    courses = [
        ('SAT Math Preparation', 'Comprehensive SAT Math prep covering algebra, geometry, and advanced topics', 'sat_prep'),
        ('SAT Reading & Writing', 'Improve reading comprehension and writing skills for SAT success', 'sat_prep'),
        ('College Application Workshop', 'Learn how to write compelling college essays and applications', 'college_readiness'),
        ('Financial Literacy', 'Understanding personal finance, budgeting, and career planning', 'life_skills')
    ]
    
    for title, desc, ctype in courses:
        course_id = append_to_sheet({
            'title': title,
            'description': desc,
            'course_type': ctype,
            'cohort_id': cohort_id
        }, SHEETS['courses'])
        course_ids.append(course_id)
    
    # Create sample assignments
    assignments = [
        ('Algebra Practice Set 1', 'Complete 20 algebra problems focusing on linear equations', course_ids[0], 7),
        ('Reading Comprehension Exercise', 'Analyze 3 reading passages and answer comprehension questions', course_ids[1], 5),
        ('Personal Statement Draft', 'Write a first draft of your college personal statement', course_ids[2], 14)
    ]
    
    for title, desc, course_id, days in assignments:
        append_to_sheet({
            'title': title,
            'description': desc,
            'course_id': course_id,
            'due_date': (datetime.now() + timedelta(days=days)).isoformat(),
            'created_at': datetime.now().isoformat()
        }, SHEETS['assignments'])
    
    # Create sample SAT test
    sample_questions = {
        "1": {
            "section": "Math",
            "type": "Multiple Choice",
            "question_text": "If 3x + 5 = 17, what is the value of x?",
            "choices": {
                "a": "2",
                "b": "4",
                "c": "6",
                "d": "8"
            },
            "correct_answer": "b"
        },
        "2": {
            "section": "Math",
            "type": "Multiple Choice",
            "question_text": "What is the area of a circle with radius 5?",
            "choices": {
                "a": "25π",
                "b": "10π",
                "c": "5π",
                "d": "15π"
            },
            "correct_answer": "a"
        },
        "3": {
            "section": "Reading",
            "type": "Multiple Choice",
            "question_text": "In the passage, the author's main argument is that:",
            "passage": "Technology has revolutionized education, making learning more accessible and engaging than ever before. However, we must be careful not to lose the human element that makes education truly meaningful.",
            "choices": {
                "a": "Technology should replace traditional teaching",
                "b": "Technology improves education but human connection remains important",
                "c": "Traditional education is superior to digital learning",
                "d": "Technology has no place in education"
            },
            "correct_answer": "b"
        }
    }
    
    test_id = append_to_sheet({
        'test_name': 'Sample SAT Practice Test',
        'questions_data': json.dumps(sample_questions),
        'time_limit': 25,
        'created_at': datetime.now().isoformat()
    }, SHEETS['sat_tests'])
    
    # Create sample test attempts
    import random
    for student_id in student_ids:
        for _ in range(random.randint(1, 3)):
            score = random.randint(400, 800)
            append_to_sheet({
                'student_id': student_id,
                'test_id': test_id,
                'answers': json.dumps({"1": "b", "2": "a", "3": "b"}),
                'score': score,
                'started_at': (datetime.now() - timedelta(days=random.randint(1, 30))).isoformat(),
                'completed_at': (datetime.now() - timedelta(days=random.randint(1, 30))).isoformat()
            }, SHEETS['test_attempts'])
    
    # Create sample survey
    survey_questions = {
        "1": {
            "question": "How confident are you in your math skills?",
            "type": "scale",
            "scale": "1-5"
        },
        "2": {
            "question": "What are your college goals?",
            "type": "text"
        },
        "3": {
            "question": "How many hours per week can you dedicate to studying?",
            "type": "multiple_choice",
            "options": ["1-3 hours", "4-6 hours", "7-10 hours", "More than 10 hours"]
        }
    }
    
    append_to_sheet({
        'title': 'Initial Learning Assessment',
        'survey_type': 'diagnostic',
        'questions': json.dumps(survey_questions),
        'cohort_id': cohort_id
    }, SHEETS['surveys'])
    
    # Create sample tutor logs
    activities = [
        "Led SAT math practice session",
        "Individual tutoring with struggling students",
        "College application workshop",
        "Parent-teacher conference calls",
        "Grading assignments and providing feedback"
    ]
    
    for i in range(10):
        append_to_sheet({
            'tutor_id': tutor_id,
            'date': (date.today() - timedelta(days=i)).isoformat(),
            'hours_worked': round(random.uniform(2, 8), 1),
            'activity_description': random.choice(activities),
            'location': 'Lincoln High School',
            'cohort_id': cohort_id,
            'created_at': datetime.now().isoformat()
        }, SHEETS['tutor_logs'])

# Main Dashboard Class
class EnrichmentDashboard(param.Parameterized):
    def __init__(self):
        super().__init__()
        self.login_view = self.create_login_view()
        self.main_view = None
        
    def create_login_view(self):
        username_input = pn.widgets.TextInput(placeholder="Username", css_classes=['input-field'])
        password_input = pn.widgets.PasswordInput(placeholder="Password", css_classes=['input-field'])
        login_button = pn.widgets.Button(name="Login", button_type="primary", css_classes=['button'])
        
        def handle_login(event):
            if session.login(username_input.value, password_input.value):
                pn.state.notifications.success(f'Welcome back, {username_input.value}!', duration=3000)
                self.main_view = self.create_main_view()
                self.param.trigger('main_view')
            else:
                pn.state.notifications.error('Invalid credentials', duration=3000)
        
        login_button.on_click(handle_login)
        
        login_card = pn.pane.HTML(f"""
        {APPLE_CSS}
        <div class="main-container">
            <div class="header">
                <h1>Enrichment Program</h1>
                <p>Empowering Students for Success</p>
            </div>
            <div class="card" style="max-width: 400px; margin: 0 auto;">
                <h2 style="margin-bottom: 30px; text-align: center;">Sign In</h2>
                <div style="margin-bottom: 20px;">
                    <label class="label">Username</label>
                    {username_input}
                </div>
                <div style="margin-bottom: 30px;">
                    <label class="label">Password</label>
                    {password_input}
                </div>
                {login_button}
                <div style="margin-top: 20px; text-align: center; color: #8E8E93;">
                    <p>Demo Credentials:</p>
                    <p>Admin: admin / admin123</p>
                    <p>Tutor: tutor1 / tutor123</p>
                    <p>Student: alice / student123</p>
                </div>
            </div>
        </div>
        """)
        
        return pn.Column(
            login_card,
            username_input,
            password_input,
            login_button,
            sizing_mode='stretch_width'
        )
    
    def create_main_view(self):
        if not session.is_authenticated():
            return self.login_view
        
        if session.role == 'student':
            return self.create_student_dashboard()
        elif session.role == 'tutor':
            return self.create_tutor_dashboard()
        elif session.role == 'admin':
            return self.create_admin_dashboard()
    
    def create_header(self, title, subtitle=""):
        logout_button = pn.widgets.Button(name="Logout", button_type="danger", css_classes=['button', 'danger'])
        
        def handle_logout(event):
            session.logout()
            self.main_view = self.login_view
            self.param.trigger('main_view')
            pn.state.notifications.info('Logged out successfully', duration=2000)
        
        logout_button.on_click(handle_logout)
        
        header_html = pn.pane.HTML(f"""
        <div class="header">
            <h1>{title}</h1>
            <p>{subtitle}</p>
        </div>
        """, sizing_mode='stretch_width')
        
        return pn.Column(
            pn.Row(header_html, logout_button, align='center'),
            sizing_mode='stretch_width'
        )
    
    def create_student_dashboard(self):
        # Get student data
        courses_df = read_sheet(SHEETS['courses'])
        student_courses = courses_df[courses_df['cohort_id'] == session.user['cohort_id']]
        
        assignments_df = read_sheet(SHEETS['assignments'])
        course_ids = student_courses['id'].tolist()
        student_assignments = assignments_df[assignments_df['course_id'].isin(course_ids)]
        
        attempts_df = read_sheet(SHEETS['test_attempts'])
        student_attempts = attempts_df[attempts_df['student_id'] == session.user_id]
        
        # Dashboard tabs
        tabs = pn.Tabs(
            ('Overview', self.create_student_overview(student_courses, student_assignments, student_attempts)),
            ('Courses', self.create_courses_view(student_courses)),
            ('SAT Tests', self.create_sat_tests_view()),
            ('Assignments', self.create_assignments_view(student_assignments)),
            ('Surveys', self.create_surveys_view()),
            tabs_location='above'
        )
        
        return pn.template.FastListTemplate(
            title="Student Dashboard",
            header=self.create_header(f"Welcome, {session.user['username']}!", "Your Learning Journey"),
            main=[tabs],
            header_background='#007AFF'
        )
    
    def create_student_overview(self, courses, assignments, attempts):
        # Calculate statistics
        total_courses = len(courses)
        pending_assignments = len(assignments[pd.to_datetime(assignments['due_date']) > datetime.now()])
        avg_score = attempts['score'].mean() if len(attempts) > 0 else 0
        
        stats_html = f"""
        {APPLE_CSS}
        <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(250px, 1fr)); gap: 20px; margin-bottom: 30px;">
            <div class="stat-card blue">
                <div class="stat-label">Enrolled Courses</div>
                <div class="stat-value">{total_courses}</div>
            </div>
            <div class="stat-card green">
                <div class="stat-label">Pending Assignments</div>
                <div class="stat-value">{pending_assignments}</div>
            </div>
            <div class="stat-card purple">
                <div class="stat-label">Average SAT Score</div>
                <div class="stat-value">{int(avg_score)}</div>
            </div>
        </div>
        """
        
        # Recent activity
        recent_html = """
        <div class="card">
            <h3>Recent Activity</h3>
            <div class="empty-state">
                <h3>📚</h3>
                <p>Your recent activities will appear here</p>
            </div>
        </div>
        """
        
        return pn.Column(
            pn.pane.HTML(stats_html),
            pn.pane.HTML(recent_html)
        )
    
    def create_courses_view(self, courses):
        if len(courses) == 0:
            return pn.pane.HTML("""
                <div class="card">
                    <div class="empty-state">
                        <h3>No Courses Yet</h3>
                        <p>You're not enrolled in any courses</p>
                    </div>
                </div>
            """)
        
        courses_html = f"""{APPLE_CSS}<div class="card">"""
        for _, course in courses.iterrows():
            courses_html += f"""
            <div style="padding: 20px; border-bottom: 1px solid #E5E5EA;">
                <h3>{course['title']}</h3>
                <p>{course['description']}</p>
                <span class="button secondary" style="margin-top: 10px;">
                    {course['course_type'].replace('_', ' ').title()}
                </span>
            </div>
            """
        courses_html += "</div>"
        
        return pn.pane.HTML(courses_html)
    
    def create_sat_tests_view(self):
        tests_df = read_sheet(SHEETS['sat_tests'])
        attempts_df = read_sheet(SHEETS['test_attempts'])
        student_attempts = attempts_df[attempts_df['student_id'] == session.user_id]
        
        # Test selection
        test_select = pn.widgets.Select(
            name='Select Test',
            options={row['test_name']: row['id'] for _, row in tests_df.iterrows()},
            css_classes=['input-field']
        )
        
        take_test_button = pn.widgets.Button(name="Take Test", button_type="primary", css_classes=['button'])
        
        def handle_take_test(event):
            if test_select.value:
                pn.state.notifications.info('Test functionality would open here', duration=3000)
        
        take_test_button.on_click(handle_take_test)
        
        # Previous attempts
        attempts_html = f"""{APPLE_CSS}<div class="card">
            <h3>Your Test History</h3>
            <div class="table-container">
                <table>
                    <thead>
                        <tr>
                            <th>Date</th>
                            <th>Test</th>
                            <th>Score</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        if len(student_attempts) > 0:
            for _, attempt in student_attempts.iterrows():
                test_name = tests_df[tests_df['id'] == attempt['test_id']]['test_name'].iloc[0]
                date_str = pd.to_datetime(attempt['completed_at']).strftime('%Y-%m-%d')
                attempts_html += f"""
                    <tr>
                        <td>{date_str}</td>
                        <td>{test_name}</td>
                        <td><strong>{attempt['score']}</strong></td>
                    </tr>
                """
        else:
            attempts_html += """
                <tr>
                    <td colspan="3" style="text-align: center; color: #8E8E93;">
                        No tests taken yet
                    </td>
                </tr>
            """
        
        attempts_html += """
                    </tbody>
                </table>
            </div>
        </div>
        """
        
        return pn.Column(
            pn.Row(test_select, take_test_button, align='center'),
            pn.pane.HTML(attempts_html)
        )
    
    def create_assignments_view(self, assignments):
        if len(assignments) == 0:
            return pn.pane.HTML("""
                <div class="card">
                    <div class="empty-state">
                        <h3>No Assignments</h3>
                        <p>You don't have any assignments yet</p>
                    </div>
                </div>
            """)
        
        assignments_html = f"""{APPLE_CSS}<div class="card">
            <h3>Your Assignments</h3>
            <div class="table-container">
                <table>
                    <thead>
                        <tr>
                            <th>Title</th>
                            <th>Description</th>
                            <th>Due Date</th>
                            <th>Status</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        for _, assignment in assignments.iterrows():
            due_date = pd.to_datetime(assignment['due_date'])
            date_str = due_date.strftime('%Y-%m-%d')
            is_overdue = due_date < datetime.now()
            status = '<span style="color: #FF3B30;">Overdue</span>' if is_overdue else '<span style="color: #34C759;">Pending</span>'
            
            assignments_html += f"""
                <tr>
                    <td><strong>{assignment['title']}</strong></td>
                    <td>{assignment['description'][:50]}...</td>
                    <td>{date_str}</td>
                    <td>{status}</td>
                </tr>
            """
        
        assignments_html += """
                    </tbody>
                </table>
            </div>
        </div>
        """
        
        return pn.pane.HTML(assignments_html)
    
    def create_surveys_view(self):
        surveys_df = read_sheet(SHEETS['surveys'])
        responses_df = read_sheet(SHEETS['survey_responses'])
        
        cohort_surveys = surveys_df[surveys_df['cohort_id'] == session.user['cohort_id']]
        student_responses = responses_df[responses_df['student_id'] == session.user_id]
        
        surveys_html = f"""{APPLE_CSS}<div class="card">
            <h3>Available Surveys</h3>
        """
        
        if len(cohort_surveys) > 0:
            for _, survey in cohort_surveys.iterrows():
                completed = len(student_responses[student_responses['survey_id'] == survey['id']]) > 0
                status = "Completed ✓" if completed else "Take Survey"
                button_class = "secondary" if completed else ""
                
                surveys_html += f"""
                <div style="padding: 20px; border-bottom: 1px solid #E5E5EA;">
                    <h4>{survey['title']}</h4>
                    <p>Type: {survey['survey_type'].replace('_', ' ').title()}</p>
                    <button class="button {button_class}">{status}</button>
                </div>
                """
        else:
            surveys_html += """
                <div class="empty-state">
                    <p>No surveys available</p>
                </div>
            """
        
        surveys_html += "</div>"
        
        return pn.pane.HTML(surveys_html)
    
    def create_tutor_dashboard(self):
        # Get tutor data
        cohorts_df = read_sheet(SHEETS['cohorts'])
        tutor_cohorts = cohorts_df[cohorts_df['tutor_id'] == session.user_id]
        
        users_df = read_sheet(SHEETS['users'])
        cohort_ids = tutor_cohorts['id'].tolist()
        students = users_df[(users_df['role'] == 'student') & (users_df['cohort_id'].isin(cohort_ids))]
        
        logs_df = read_sheet(SHEETS['tutor_logs'])
        tutor_logs = logs_df[logs_df['tutor_id'] == session.user_id]
        
        # Dashboard tabs
        tabs = pn.Tabs(
            ('Overview', self.create_tutor_overview(tutor_cohorts, students, tutor_logs)),
            ('Students', self.create_students_management(students)),
            ('Log Hours', self.create_log_hours(tutor_cohorts)),
            ('Reports', self.create_tutor_reports(students)),
            tabs_location='above'
        )
        
        return pn.template.FastListTemplate(
            title="Tutor Dashboard",
            header=self.create_header(f"Welcome, {session.user['username']}!", "Tutor Portal"),
            main=[tabs],
            header_background='#34C759'
        )
    
    def create_tutor_overview(self, cohorts, students, logs):
        total_students = len(students)
        total_cohorts = len(cohorts)
        total_hours = logs['hours_worked'].sum() if len(logs) > 0 else 0
        
        stats_html = f"""
        {APPLE_CSS}
        <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(250px, 1fr)); gap: 20px; margin-bottom: 30px;">
            <div class="stat-card blue">
                <div class="stat-label">Active Cohorts</div>
                <div class="stat-value">{total_cohorts}</div>
            </div>
            <div class="stat-card green">
                <div class="stat-label">Total Students</div>
                <div class="stat-value">{total_students}</div>
            </div>
            <div class="stat-card purple">
                <div class="stat-label">Hours Logged</div>
                <div class="stat-value">{total_hours:.1f}</div>
            </div>
        </div>
        """
        
        # Recent logs
        recent_logs_html = """<div class="card"><h3>Recent Activity Logs</h3>"""
        
        if len(logs) > 0:
            recent_logs = logs.nlargest(5, 'created_at')
            recent_logs_html += """
                <div class="table-container">
                    <table>
                        <thead>
                            <tr>
                                <th>Date</th>
                                <th>Hours</th>
                                <th>Activity</th>
                            </tr>
                        </thead>
                        <tbody>
            """
            
            for _, log in recent_logs.iterrows():
                date_str = pd.to_datetime(log['date']).strftime('%Y-%m-%d')
                recent_logs_html += f"""
                    <tr>
                        <td>{date_str}</td>
                        <td>{log['hours_worked']}</td>
                        <td>{log['activity_description'][:50]}...</td>
                    </tr>
                """
            
            recent_logs_html += """
                        </tbody>
                    </table>
                </div>
            """
        else:
            recent_logs_html += """
                <div class="empty-state">
                    <p>No activity logs yet</p>
                </div>
            """
        
        recent_logs_html += "</div>"
        
        return pn.Column(
            pn.pane.HTML(stats_html),
            pn.pane.HTML(recent_logs_html)
        )
    
    def create_students_management(self, students):
        if len(students) == 0:
            return pn.pane.HTML("""
                <div class="card">
                    <div class="empty-state">
                        <h3>No Students</h3>
                        <p>No students assigned to your cohorts yet</p>
                    </div>
                </div>
            """)
        
        # Get test attempts for progress
        attempts_df = read_sheet(SHEETS['test_attempts'])
        
        students_html = f"""{APPLE_CSS}<div class="card">
            <h3>Your Students</h3>
            <div class="table-container">
                <table>
                    <thead>
                        <tr>
                            <th>Name</th>
                            <th>Email</th>
                            <th>Tests Taken</th>
                            <th>Avg Score</th>
                            <th>Actions</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        for _, student in students.iterrows():
            student_attempts = attempts_df[attempts_df['student_id'] == student['id']]
            tests_taken = len(student_attempts)
            avg_score = student_attempts['score'].mean() if tests_taken > 0 else 0
            
            students_html += f"""
                <tr>
                    <td><strong>{student['username']}</strong></td>
                    <td>{student['email']}</td>
                    <td>{tests_taken}</td>
                    <td>{int(avg_score)}</td>
                    <td>
                        <button class="button secondary" style="padding: 8px 16px; font-size: 14px;">
                            View Progress
                        </button>
                    </td>
                </tr>
            """
        
        students_html += """
                    </tbody>
                </table>
            </div>
        </div>
        """
        
        return pn.pane.HTML(students_html)
    
    def create_log_hours(self, cohorts):
        # Form inputs
        date_input = pn.widgets.DatePicker(name="Date", value=date.today(), css_classes=['input-field'])
        hours_input = pn.widgets.FloatInput(name="Hours Worked", value=0, step=0.5, bounds=(0, 24), css_classes=['input-field'])
        cohort_select = pn.widgets.Select(
            name="Cohort",
            options={row['name']: row['id'] for _, row in cohorts.iterrows()},
            css_classes=['input-field']
        )
        location_input = pn.widgets.TextInput(name="Location", placeholder="e.g., Lincoln High School", css_classes=['input-field'])
        activity_input = pn.widgets.TextAreaInput(
            name="Activity Description",
            placeholder="Describe your tutoring activities...",
            height=100,
            css_classes=['input-field']
        )
        
        submit_button = pn.widgets.Button(name="Log Hours", button_type="primary", css_classes=['button'])
        
        def handle_submit(event):
            if hours_input.value > 0 and activity_input.value and cohort_select.value:
                append_to_sheet({
                    'tutor_id': session.user_id,
                    'date': date_input.value.isoformat(),
                    'hours_worked': hours_input.value,
                    'activity_description': activity_input.value,
                    'location': location_input.value,
                    'cohort_id': cohort_select.value,
                    'created_at': datetime.now().isoformat()
                }, SHEETS['tutor_logs'])
                
                pn.state.notifications.success('Hours logged successfully!', duration=3000)
                
                # Reset form
                hours_input.value = 0
                activity_input.value = ""
                location_input.value = ""
            else:
                pn.state.notifications.error('Please fill all required fields', duration=3000)
        
        submit_button.on_click(handle_submit)
        
        form_html = f"""{APPLE_CSS}
        <div class="card">
            <h3>Log Tutoring Hours</h3>
            <p style="color: #8E8E93; margin-bottom: 30px;">
                Record your tutoring activities and hours worked
            </p>
        </div>
        """
        
        return pn.Column(
            pn.pane.HTML(form_html),
            date_input,
            hours_input,
            cohort_select,
            location_input,
            activity_input,
            submit_button,
            width=600
        )
    
    def create_tutor_reports(self, students):
        # Summary statistics
        attempts_df = read_sheet(SHEETS['test_attempts'])
        logs_df = read_sheet(SHEETS['tutor_logs'])
        tutor_logs = logs_df[logs_df['tutor_id'] == session.user_id]
        
        # Calculate metrics
        student_ids = students['id'].tolist()
        all_attempts = attempts_df[attempts_df['student_id'].isin(student_ids)]
        
        avg_score = all_attempts['score'].mean() if len(all_attempts) > 0 else 0
        total_tests = len(all_attempts)
        improvement_rate = 0  # Would calculate based on date ordering
        
        # This month's hours
        this_month = pd.to_datetime(tutor_logs['date']).dt.to_period('M') == pd.Period(date.today(), 'M')
        monthly_hours = tutor_logs.loc[this_month, 'hours_worked'].sum() if any(this_month) else 0
        
        report_html = f"""
        {APPLE_CSS}
        <div class="card">
            <h3>Performance Report</h3>
            <div style="display: grid; grid-template-columns: repeat(2, 1fr); gap: 30px; margin-top: 30px;">
                <div>
                    <h4>Student Performance</h4>
                    <p><strong>Average SAT Score:</strong> {int(avg_score)}</p>
                    <p><strong>Total Tests Taken:</strong> {total_tests}</p>
                    <p><strong>Improvement Rate:</strong> +{improvement_rate}%</p>
                </div>
                <div>
                    <h4>Your Activity</h4>
                    <p><strong>Hours This Month:</strong> {monthly_hours:.1f}</p>
                    <p><strong>Total Hours:</strong> {tutor_logs['hours_worked'].sum():.1f}</p>
                    <p><strong>Active Students:</strong> {len(students)}</p>
                </div>
            </div>
        </div>
        """
        
        return pn.pane.HTML(report_html)
    
    def create_admin_dashboard(self):
        # Get all data for admin overview
        users_df = read_sheet(SHEETS['users'])
        schools_df = read_sheet(SHEETS['schools'])
        cohorts_df = read_sheet(SHEETS['cohorts'])
        attempts_df = read_sheet(SHEETS['test_attempts'])
        logs_df = read_sheet(SHEETS['tutor_logs'])
        
        # Dashboard tabs
        tabs = pn.Tabs(
            ('Overview', self.create_admin_overview(users_df, schools_df, cohorts_df, attempts_df, logs_df)),
            ('Users', self.create_users_management(users_df)),
            ('Schools', self.create_schools_management(schools_df)),
            ('Reports', self.create_admin_reports(users_df, attempts_df, logs_df)),
            tabs_location='above'
        )
        
        return pn.template.FastListTemplate(
            title="Admin Dashboard",
            header=self.create_header("Admin Portal", "System Management"),
            main=[tabs],
            header_background='#5856D6'
        )
    
    def create_admin_overview(self, users_df, schools_df, cohorts_df, attempts_df, logs_df):
        total_students = len(users_df[users_df['role'] == 'student'])
        total_tutors = len(users_df[users_df['role'] == 'tutor'])
        total_schools = len(schools_df)
        total_cohorts = len(cohorts_df)
        
        stats_html = f"""
        {APPLE_CSS}
        <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(250px, 1fr)); gap: 20px; margin-bottom: 30px;">
            <div class="stat-card">
                <div class="stat-label">Total Students</div>
                <div class="stat-value">{total_students}</div>
            </div>
            <div class="stat-card blue">
                <div class="stat-label">Total Tutors</div>
                <div class="stat-value">{total_tutors}</div>
            </div>
            <div class="stat-card green">
                <div class="stat-label">Schools</div>
                <div class="stat-value">{total_schools}</div>
            </div>
            <div class="stat-card purple">
                <div class="stat-label">Active Cohorts</div>
                <div class="stat-value">{total_cohorts}</div>
            </div>
        </div>
        """
        
        # Recent activity
        recent_html = """<div class="card"><h3>Recent System Activity</h3>"""
        
        # Recent test attempts
        if len(attempts_df) > 0:
            recent_attempts = attempts_df.nlargest(5, 'completed_at')
            recent_html += """
                <h4>Recent Tests</h4>
                <div class="table-container">
                    <table>
                        <thead>
                            <tr>
                                <th>Student ID</th>
                                <th>Score</th>
                                <th>Date</th>
                            </tr>
                        </thead>
                        <tbody>
            """
            
            for _, attempt in recent_attempts.iterrows():
                date_str = pd.to_datetime(attempt['completed_at']).strftime('%Y-%m-%d %H:%M')
                recent_html += f"""
                    <tr>
                        <td>Student #{attempt['student_id']}</td>
                        <td>{attempt['score']}</td>
                        <td>{date_str}</td>
                    </tr>
                """
            
            recent_html += """
                        </tbody>
                    </table>
                </div>
            """
        
        recent_html += "</div>"
        
        return pn.Column(
            pn.pane.HTML(stats_html),
            pn.pane.HTML(recent_html)
        )
    
    def create_users_management(self, users_df):
        users_html = f"""{APPLE_CSS}<div class="card">
            <h3>User Management</h3>
            <div class="table-container">
                <table>
                    <thead>
                        <tr>
                            <th>ID</th>
                            <th>Username</th>
                            <th>Email</th>
                            <th>Role</th>
                            <th>Status</th>
                            <th>Actions</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        for _, user in users_df.iterrows():
            status = "Active" if user['is_active'] else "Inactive"
            status_color = "#34C759" if user['is_active'] else "#FF3B30"
            
            users_html += f"""
                <tr>
                    <td>{user['id']}</td>
                    <td><strong>{user['username']}</strong></td>
                    <td>{user['email']}</td>
                    <td>{user['role'].title()}</td>
                    <td><span style="color: {status_color};">{status}</span></td>
                    <td>
                        <button class="button secondary" style="padding: 6px 12px; font-size: 13px;">
                            Edit
                        </button>
                    </td>
                </tr>
            """
        
        users_html += """
                    </tbody>
                </table>
            </div>
        </div>
        """
        
        return pn.pane.HTML(users_html)
    
    def create_schools_management(self, schools_df):
        schools_html = f"""{APPLE_CSS}<div class="card">
            <h3>School Management</h3>
            <button class="button" style="margin-bottom: 20px;">Add New School</button>
            <div class="table-container">
                <table>
                    <thead>
                        <tr>
                            <th>ID</th>
                            <th>Name</th>
                            <th>Address</th>
                            <th>Contact Email</th>
                            <th>Actions</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        for _, school in schools_df.iterrows():
            schools_html += f"""
                <tr>
                    <td>{school['id']}</td>
                    <td><strong>{school['name']}</strong></td>
                    <td>{school['address']}</td>
                    <td>{school['contact_email']}</td>
                    <td>
                        <button class="button secondary" style="padding: 6px 12px; font-size: 13px;">
                            Edit
                        </button>
                    </td>
                </tr>
            """
        
        schools_html += """
                    </tbody>
                </table>
            </div>
        </div>
        """
        
        return pn.pane.HTML(schools_html)
    
    def create_admin_reports(self, users_df, attempts_df, logs_df):
        # Calculate system-wide metrics
        total_students = len(users_df[users_df['role'] == 'student'])
        total_tests = len(attempts_df)
        avg_system_score = attempts_df['score'].mean() if total_tests > 0 else 0
        total_tutor_hours = logs_df['hours_worked'].sum()
        
        # Score distribution
        score_bins = pd.cut(attempts_df['score'], bins=[0, 400, 500, 600, 700, 800], labels=['<400', '400-500', '500-600', '600-700', '700-800'])
        score_dist = score_bins.value_counts().sort_index()
        
        report_html = f"""
        {APPLE_CSS}
        <div class="card">
            <h3>System Reports</h3>
            
            <div style="display: grid; grid-template-columns: repeat(2, 1fr); gap: 40px; margin-top: 30px;">
                <div>
                    <h4>Overall Performance</h4>
                    <p><strong>Total Students:</strong> {total_students}</p>
                    <p><strong>Total Tests Taken:</strong> {total_tests}</p>
                    <p><strong>System Average Score:</strong> {int(avg_system_score)}</p>
                    <p><strong>Total Tutor Hours:</strong> {total_tutor_hours:.1f}</p>
                </div>
                
                <div>
                    <h4>Score Distribution</h4>
        """
        
        if len(score_dist) > 0:
            for range_label, count in score_dist.items():
                percentage = (count / total_tests * 100) if total_tests > 0 else 0
                report_html += f"""
                    <div style="margin-bottom: 10px;">
                        <div style="display: flex; justify-content: space-between; margin-bottom: 5px;">
                            <span>{range_label}</span>
                            <span>{count} ({percentage:.1f}%)</span>
                        </div>
                        <div class="progress-bar">
                            <div class="progress-fill" style="width: {percentage}%"></div>
                        </div>
                    </div>
                """
        
        report_html += """
                </div>
            </div>
            
            <div style="margin-top: 40px;">
                <button class="button">Export Full Report</button>
                <button class="button secondary">Download Excel Data</button>
            </div>
        </div>
        """
        
        return pn.pane.HTML(report_html)
    
    @param.depends('main_view')
    def view(self):
        return self.main_view or self.login_view

# Initialize and run
init_excel()
dashboard = EnrichmentDashboard()

# Create the app
template = pn.template.FastListTemplate(
    title="Enrichment Program Dashboard",
    sidebar=[],
    main=[dashboard.view],
    header_background='#007AFF'
)

# Serve the app
if __name__ == "__main__":
    print("🎓 Enrichment Program Dashboard Starting...")
    print("📊 Data stored in: enrichment_data.xlsx")
    print("🌐 Access at: http://localhost:5006")
    print("\n📱 Login Credentials:")
    print("   Admin: admin / admin123")
    print("   Tutor: tutor1 / tutor123")
    print("   Student: alice / student123 (or bob, carol)")
    
    template.servable()
    pn.serve(template, port=5006, show=True)